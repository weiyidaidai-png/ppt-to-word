#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
创建一个示例 Word 文档用于测试
"""

from docx import Document
from docx.shared import Pt


def create_example_word(doc_path: str) -> None:
    """
    创建示例 Word 文档，包含标题和正文内容
    """
    doc = Document()

    # 添加标题
    title = doc.add_heading("项目报告", 0)
    run = title.runs[0]
    run.font.size = Pt(24)

    # 添加正文
    doc.add_paragraph("本报告总结了最近项目的进展情况。")
    doc.add_paragraph()

    # 添加子标题 1
    subtitle1 = doc.add_heading("项目概述", level=1)
    subtitle1.runs[0].font.size = Pt(16)

    doc.add_paragraph("该项目旨在开发一套自动化办公工具。")
    doc.add_paragraph("主要功能包括：")
    doc.add_paragraph("• 文档格式转换", style="List Bullet")
    doc.add_paragraph("• 数据提取与分析", style="List Bullet")
    doc.add_paragraph("• 自动化报告生成", style="List Bullet")

    # 添加子标题 2
    subtitle2 = doc.add_heading("实施进度", level=1)
    subtitle2.runs[0].font.size = Pt(16)

    doc.add_paragraph("目前项目已完成 75%，主要完成的工作包括：")
    doc.add_paragraph("1. 核心功能开发完成", style="List Number")
    doc.add_paragraph("2. 用户界面设计", style="List Number")
    doc.add_paragraph("3. 基础测试通过", style="List Number")

    # 添加子标题 3
    doc.add_heading("下一步计划", level=1)

    doc.add_paragraph("接下来需要完成：")
    doc.add_paragraph("• 系统集成测试")
    doc.add_paragraph("• 性能优化")
    doc.add_paragraph("• 用户培训文档准备")

    # 保存文档
    doc.save(doc_path)
    print(f"示例 Word 文档已创建在: {doc_path}")


if __name__ == "__main__":
    import os
    example_path = os.path.join(os.path.dirname(__file__), "example.docx")
    create_example_word(example_path)
    print(f"请运行以下命令测试转换:")
    print(f"python word_to_ppt.py \"{example_path}\" \"example.pptx\"")
