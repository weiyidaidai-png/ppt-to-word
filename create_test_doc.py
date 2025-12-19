#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def create_test_document():
    doc = Document()

    # 添加测试标题和内容
    doc.add_heading('一级标题 - 主要章节', level=1)
    p = doc.add_paragraph('这是一级标题下面的正文内容，应该被正确识别并归属到一级标题下。')

    doc.add_heading('二级标题 - 子章节', level=2)
    p = doc.add_paragraph('这是二级标题的内容。')

    doc.add_heading('三级标题 - 详细内容', level=3)
    p = doc.add_paragraph('三级标题的正文。')

    doc.add_heading('四级标题 - 更详细', level=4)
    p = doc.add_paragraph('四级标题内容。')

    doc.add_heading('五级标题 - 进一步', level=5)
    p = doc.add_paragraph('五级标题内容。')

    doc.add_heading('六级标题 - 深入', level=6)
    p = doc.add_paragraph('六级标题内容。')

    doc.add_heading('七级标题 - 最底层', level=7)
    p = doc.add_paragraph('七级标题内容。这是最深层级的标题。')

    doc.save('C:/Users/18167/code/ppt to word/test_multilevel_headings.docx')
    print("测试文档已创建在: C:/Users/18167/code/ppt to word/test_multilevel_headings.docx")

if __name__ == '__main__':
    create_test_document()
